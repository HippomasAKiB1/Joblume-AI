"""PDF and DOCX resume generator using reportlab and python-docx."""
import io
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch, mm
from reportlab.lib.colors import HexColor
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, HRFlowable
from reportlab.lib.enums import TA_LEFT, TA_CENTER


PRIMARY_COLOR = HexColor("#4F46E5")
TEXT_COLOR = HexColor("#1E293B")
SUBTLE_COLOR = HexColor("#64748B")


def _get_styles():
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle("ResumeName", fontSize=20, leading=24, textColor=PRIMARY_COLOR, alignment=TA_CENTER, spaceAfter=4, fontName="Helvetica-Bold"))
    styles.add(ParagraphStyle("ResumeContact", fontSize=9, leading=12, textColor=SUBTLE_COLOR, alignment=TA_CENTER, spaceAfter=12))
    styles.add(ParagraphStyle("SectionTitle", fontSize=12, leading=15, textColor=PRIMARY_COLOR, spaceAfter=4, spaceBefore=12, fontName="Helvetica-Bold"))
    styles.add(ParagraphStyle("ItemTitle", fontSize=10, leading=13, textColor=TEXT_COLOR, fontName="Helvetica-Bold"))
    styles.add(ParagraphStyle("ItemSubtitle", fontSize=9, leading=12, textColor=SUBTLE_COLOR))
    styles.add(ParagraphStyle("BodyText2", fontSize=9.5, leading=13, textColor=TEXT_COLOR))
    styles.add(ParagraphStyle("BulletItem", fontSize=9.5, leading=13, textColor=TEXT_COLOR, leftIndent=12, bulletIndent=0))
    return styles


def generate_pdf(resume_data: dict) -> bytes:
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4, leftMargin=0.6*inch, rightMargin=0.6*inch, topMargin=0.5*inch, bottomMargin=0.5*inch)
    styles = _get_styles()
    story = []
    pi = resume_data.get("personalInfo", {})
    story.append(Paragraph(pi.get("fullName", "Your Name"), styles["ResumeName"]))
    contacts = " | ".join(filter(None, [pi.get("email"), pi.get("phone"), pi.get("location"), pi.get("linkedin")]))
    if contacts:
        story.append(Paragraph(contacts, styles["ResumeContact"]))
    story.append(HRFlowable(width="100%", thickness=1, color=PRIMARY_COLOR, spaceAfter=8))

    summary = resume_data.get("summary", "")
    if summary:
        story.append(Paragraph("PROFESSIONAL SUMMARY", styles["SectionTitle"]))
        story.append(Paragraph(summary, styles["BodyText2"]))

    for exp in resume_data.get("experience", []):
        story.append(Paragraph("EXPERIENCE", styles["SectionTitle"])) if resume_data.get("experience", []).index(exp) == 0 else None
        title_line = f"{exp.get('title', '')} — {exp.get('company', '')}"
        story.append(Paragraph(title_line, styles["ItemTitle"]))
        date_line = f"{exp.get('startDate', '')} – {exp.get('endDate', 'Present')}"
        story.append(Paragraph(date_line, styles["ItemSubtitle"]))
        for bullet in exp.get("bullets", []):
            story.append(Paragraph(f"• {bullet}", styles["BulletItem"]))
        story.append(Spacer(1, 4))

    education = resume_data.get("education", [])
    if education:
        story.append(Paragraph("EDUCATION", styles["SectionTitle"]))
        for edu in education:
            story.append(Paragraph(f"{edu.get('degree', '')} — {edu.get('institution', '')}", styles["ItemTitle"]))
            story.append(Paragraph(f"{edu.get('startDate', '')} – {edu.get('endDate', '')}", styles["ItemSubtitle"]))

    skills = resume_data.get("skills", [])
    if skills:
        story.append(Paragraph("SKILLS", styles["SectionTitle"]))
        story.append(Paragraph(", ".join(skills), styles["BodyText2"]))

    projects = resume_data.get("projects", [])
    if projects:
        story.append(Paragraph("PROJECTS", styles["SectionTitle"]))
        for proj in projects:
            story.append(Paragraph(proj.get("name", ""), styles["ItemTitle"]))
            story.append(Paragraph(proj.get("description", ""), styles["BodyText2"]))

    certs = resume_data.get("certifications", [])
    if certs:
        story.append(Paragraph("CERTIFICATIONS", styles["SectionTitle"]))
        for c in certs:
            story.append(Paragraph(f"{c.get('name', '')} — {c.get('issuer', '')} ({c.get('date', '')})", styles["BodyText2"]))

    # Filter out None values from story
    story = [s for s in story if s is not None]
    doc.build(story)
    return buffer.getvalue()


def generate_docx(resume_data: dict) -> bytes:
    from docx import Document
    from docx.shared import Pt, RGBColor, Inches
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()
    style = doc.styles["Normal"]
    style.font.size = Pt(10)
    style.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)

    pi = resume_data.get("personalInfo", {})
    name_para = doc.add_heading(pi.get("fullName", "Your Name"), level=0)
    name_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    contacts = " | ".join(filter(None, [pi.get("email"), pi.get("phone"), pi.get("location")]))
    if contacts:
        cp = doc.add_paragraph(contacts)
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER

    summary = resume_data.get("summary", "")
    if summary:
        doc.add_heading("Professional Summary", level=1)
        doc.add_paragraph(summary)

    experience = resume_data.get("experience", [])
    if experience:
        doc.add_heading("Experience", level=1)
        for exp in experience:
            doc.add_heading(f"{exp.get('title', '')} — {exp.get('company', '')}", level=2)
            doc.add_paragraph(f"{exp.get('startDate', '')} – {exp.get('endDate', 'Present')}")
            for bullet in exp.get("bullets", []):
                doc.add_paragraph(bullet, style="List Bullet")

    education = resume_data.get("education", [])
    if education:
        doc.add_heading("Education", level=1)
        for edu in education:
            doc.add_paragraph(f"{edu.get('degree', '')} — {edu.get('institution', '')} ({edu.get('startDate', '')} – {edu.get('endDate', '')})")

    skills = resume_data.get("skills", [])
    if skills:
        doc.add_heading("Skills", level=1)
        doc.add_paragraph(", ".join(skills))

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
